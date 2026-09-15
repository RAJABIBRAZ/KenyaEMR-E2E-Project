#!/usr/bin/env python3
"""Inventory a local TaifaCare test workbook without copying clinical test data.

The source workbook is read-only. The generated CSV/XLSX contain case metadata,
traceability, conservative execution-safety classifications, and coverage status.
"""

from __future__ import annotations

import argparse
import csv
import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


CASE_ID = re.compile(r"^(?:TC\d+|FT-\d+)$")
EDGE_MARKER = re.compile(
    r"\b(?:invalid|empty|missing|duplicate|timeout|boundary|negative|unauthorized|out of stock|"
    r"expired|future|past|cancel|offline|partial|error|fail|not found|no data)\b|[<>]=?",
    re.IGNORECASE,
)
STATE_CHANGE = re.compile(
    r"\b(?:register|create|save|record|dispens\w*|prescrib\w*|admit|assign|allocate|"
    r"transfer|discharg\w*|schedule|upload|delete|edit|update|mark|collect|call patient|"
    r"send otp|reset password|grant|revoke|enrol\w*|start visit|submit|issue|payment|"
    r"invoice|order)\b",
    re.IGNORECASE,
)
GENERIC_EXPECTED = re.compile(
    r"\b(?:works as expected|functions correctly|works correctly|proper validation|"
    r"as per mockup|proper workflow|works with proper|works as expected|functions with proper)\b",
    re.IGNORECASE,
)
GENERIC_STEPS = re.compile(
    r"\b(?:perform|test|verify functionality|verify workflow|as per mockup)\b",
    re.IGNORECASE,
)
SENSITIVE_MODULES = {
    "04_Triage",
    "05_Consultation",
    "06_Patient_Chart",
    "07_Laboratory",
    "08_Pharmacy",
    "09_Admissions",
    "10_MCH",
    "11_Imaging",
    "12_Procedures",
    "13_Billing",
    "14_Reports",
    "16_Partograph",
    "17_EHR_Reports",
    "18_HIV_Reports",
}
HIGH_IMPACT_MODULES = {
    "04_Triage",
    "05_Consultation",
    "07_Laboratory",
    "08_Pharmacy",
    "09_Admissions",
    "10_MCH",
    "11_Imaging",
    "12_Procedures",
    "13_Billing",
    "15_System",
    "16_Partograph",
}
SPEC_NAMES = {
    "01_Login": "login",
    "02_Dashboard": "dashboard",
    "03_Registration": "registration",
    "04_Triage": "triage",
    "05_Consultation": "consultation",
    "06_Patient_Chart": "patient-chart",
    "07_Laboratory": "laboratory",
    "08_Pharmacy": "pharmacy",
    "09_Admissions": "admissions",
    "10_MCH": "mch",
    "11_Imaging": "imaging",
    "12_Procedures": "procedures",
    "13_Billing": "billing",
    "14_Reports": "reports",
    "15_System": "system",
    "16_Partograph": "partograph",
    "17_EHR_Reports": "ehr-reports",
    "18_HIV_Reports": "hiv-reports",
}
INVENTORY_COLUMNS = [
    "Source Key",
    "Workbook TC ID",
    "ID Occurrences",
    "Source Sheet",
    "Source Row",
    "Module",
    "Sub-Module",
    "Test Case Name",
    "Priority",
    "Workbook Test Status",
    "Workbook Dev Status",
    "Case Type",
    "Execution Safety",
    "Specification Flags",
    "Existing Playwright Coverage",
    "Automation State",
    "Suggested Playwright Spec",
    "Fixture / Approval Needed",
]


def value(cell) -> str:
    return str(cell.value).strip() if cell.value is not None else ""


def extract_cases(source: Path) -> list[dict[str, str]]:
    workbook = load_workbook(source, data_only=True)
    cases: list[dict[str, str]] = []
    for sheet in workbook.worksheets:
        if sheet.title == "Summary":
            continue
        for row in sheet.iter_rows(min_row=2, max_col=15):
            case_id = value(row[0])
            if not CASE_ID.fullmatch(case_id):
                continue
            cases.append(
                {
                    "source_key": f"{sheet.title}!R{row[0].row}",
                    "id": case_id,
                    "sheet": sheet.title,
                    "row": str(row[0].row),
                    "module": value(row[1]),
                    "submodule": value(row[2]),
                    "name": value(row[3]),
                    "priority": value(row[5]),
                    "steps": value(row[7]),
                    "expected": value(row[8]),
                    "test_status": value(row[10]),
                    "dev_status": value(row[12]),
                }
            )
    return cases


def safety_class(case: dict[str, str]) -> tuple[str, str]:
    text = f"{case['name']}\n{case['steps']}"
    if case["sheet"] == "16_Partograph":
        return "Clinical-state change possible", "Approved synthetic labor case; clinical threshold review; cleanup"
    if STATE_CHANGE.search(text):
        if case["sheet"] in HIGH_IMPACT_MODULES or case["sheet"] == "03_Registration":
            return "High-impact write/side effect possible", "Approval; synthetic fixture; role; cleanup/idempotency"
        return "State change/side effect possible", "Approval; synthetic fixture and cleanup"
    if case["sheet"] in SENSITIVE_MODULES or case["sheet"] == "03_Registration":
        return "Sensitive read-only candidate", "De-identified seeded fixture; authorized role"
    return "Read-only UI candidate", "Current QA role and stable UI assertion"


def existing_coverage(case: dict[str, str]) -> tuple[str, str]:
    if case["sheet"] == "01_Login" and case["id"] == "TC003":
        return "Partial: e2e/specs/login.spec.ts", "Partial existing E2E; not verified on current QA"
    if case["sheet"] == "15_System" and case["id"] == "TC299":
        return "Partial: e2e/specs/logout.spec.ts", "Partial existing E2E; not verified on current QA"
    return "None", "Not automated"


def inventory_rows(cases: list[dict[str, str]]) -> list[list[str]]:
    id_count = Counter(case["id"] for case in cases)
    scenario_count = Counter(
        (case["sheet"], case["name"].casefold(), case["steps"].casefold()) for case in cases
    )
    rows: list[list[str]] = []
    for case in cases:
        flags: list[str] = []
        if id_count[case["id"]] > 1:
            flags.append("Workbook TC ID reused")
        if scenario_count[(case["sheet"], case["name"].casefold(), case["steps"].casefold())] > 1:
            flags.append("Possible repeated scenario row")
        if GENERIC_EXPECTED.search(case["expected"]):
            flags.append("Expected result needs measurable assertion")
        if GENERIC_STEPS.search(case["steps"]):
            flags.append("Steps need concrete UI action/test data")
        if not case["test_status"]:
            flags.append("Workbook test status blank")
        if not case["priority"]:
            flags.append("Priority blank")

        safety, fixture = safety_class(case)
        coverage, state = existing_coverage(case)
        case_type = "Negative/boundary candidate" if EDGE_MARKER.search(
            f"{case['name']}\n{case['steps']}"
        ) else "Functional/UI"
        rows.append(
            [
                case["source_key"],
                case["id"],
                str(id_count[case["id"]]),
                case["sheet"],
                case["row"],
                case["module"],
                case["submodule"],
                case["name"],
                case["priority"],
                case["test_status"],
                case["dev_status"],
                case_type,
                safety,
                "; ".join(flags),
                coverage,
                state,
                f"e2e/cases/{SPEC_NAMES.get(case['sheet'], 'review')}.spec.ts",
                fixture,
            ]
        )
    return rows


def write_csv(path: Path, header: list[str], rows: list[list[str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(header)
        writer.writerows(rows)


def style_sheet(sheet, header: list[str], rows: list[list[str]]) -> None:
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"
    sheet.row_dimensions[1].height = 30
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill(fill_type="solid", fgColor="174A67")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for index, name in enumerate(header, start=1):
        width = 15
        if name in {"Test Case Name", "Scenario", "Expected Result", "Specification Flags"}:
            width = 48
        elif name in {"Fixture / Approval Needed", "Steps", "Existing Playwright Coverage"}:
            width = 38
        elif name in {"Source Key", "Execution Safety", "Automation State"}:
            width = 30
        sheet.column_dimensions[get_column_letter(index)].width = width


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="Local source .xlsx workbook")
    parser.add_argument("--output-dir", type=Path, default=Path("docs"))
    args = parser.parse_args()
    cases = extract_cases(args.source)
    rows = inventory_rows(cases)
    inventory_csv = args.output_dir / "taifacare-test-case-coverage.csv"
    edge_csv = args.output_dir / "taifacare-proposed-edge-cases.csv"
    workbook_path = args.output_dir / "taifacare-test-coverage-and-edge-cases.xlsx"
    write_csv(inventory_csv, INVENTORY_COLUMNS, rows)

    with edge_csv.open(newline="", encoding="utf-8") as stream:
        edge_reader = csv.reader(stream)
        edge_header = next(edge_reader)
        edge_rows = list(edge_reader)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    ids = Counter(case["id"] for case in cases)
    module_counts = Counter(case["sheet"] for case in cases)
    safety_counts = Counter(row[12] for row in rows)
    summary_rows = [
        ["Metric", "Value"],
        ["Actual populated case rows", len(cases)],
        ["Unique workbook IDs", len(ids)],
        ["IDs reused", sum(count > 1 for count in ids.values())],
        ["Extra occurrences from reused IDs", sum(count - 1 for count in ids.values() if count > 1)],
        ["Proposed additional edge cases", len(edge_rows)],
        ["Existing Playwright coverage", "2 partial mappings; current QA not verified"],
        ["Execution policy", "No clinical/financial writes automatically run by this workbook"],
        ["Source", args.source.name],
        ["Source date (filename)", "2025-11-11"],
        ["Summary-tab count", "Conflicts: title 481; metadata 595; populated rows counted above"],
    ]
    summary_rows.extend([[f"Module: {name}", count] for name, count in sorted(module_counts.items())])
    summary_rows.extend([[f"Safety: {name}", count] for name, count in sorted(safety_counts.items())])
    for row in summary_rows:
        summary.append(row)
    summary.column_dimensions["A"].width = 44
    summary.column_dimensions["B"].width = 90
    summary.freeze_panes = "A2"
    for cell in summary[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill(fill_type="solid", fgColor="174A67")

    style_sheet(workbook.create_sheet("Source Case Inventory"), INVENTORY_COLUMNS, rows)
    style_sheet(workbook.create_sheet("Proposed Edge Cases"), edge_header, edge_rows)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    workbook.save(workbook_path)
    print(
        f"cases={len(cases)} unique_ids={len(ids)} reused_ids={sum(count > 1 for count in ids.values())} "
        f"edge_cases={len(edge_rows)}"
    )
    print(f"inventory={inventory_csv}\nworkbook={workbook_path}")


if __name__ == "__main__":
    main()
