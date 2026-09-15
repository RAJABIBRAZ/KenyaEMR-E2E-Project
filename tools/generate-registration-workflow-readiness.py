#!/usr/bin/env python3
"""Build a source-keyed Registration workflow matrix without copying test data."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SHEET = "03_Registration"
PREFIX = SHEET + "!R"
FIELDS = [
    "Source Key",
    "Workbook TC ID",
    "Canonical Source Key",
    "Exact Duplicate",
    "Workflow",
    "Scenario",
    "Execution Safety",
    "Implementation State",
    "Execution State",
    "Result",
    "Required Decision or Data",
    "Next Action",
]

REGISTRY = {3, 4, 5, 6, 7, 12, 19, 31}
OTP = {8, 9, 10, 11, 13}
PATIENT_WRITES = {14, 16, 18, 20, 21, 22, 24, 25, 34, 35, 37}
SAVE_VALIDATIONS = {15, 26, 27, 28, 29}


def read_csv_rows(path: Path, key: str) -> dict[str, dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as source:
        return {row[key]: row for row in csv.DictReader(source)}


def workflow(row_number: int) -> str:
    if row_number == 2:
        return "Registration landing navigation"
    if row_number in REGISTRY:
        return "National ID and Client Registry/HIE"
    if row_number in OTP:
        return "OTP and dependent verification"
    if row_number in PATIENT_WRITES:
        return "Create, emergency, edit, or start visit"
    return "Patient form and validation"


def dependency(row_number: int) -> str:
    if row_number == 2:
        return "Read-only UAT account"
    if row_number == 3:
        return "Current National ID control confirmed on UAT"
    if row_number in {4, 5}:
        return "Approved synthetic National ID format; no real identity"
    if row_number in REGISTRY:
        return "Approved de-identified HIE/Client Registry fixture and read role"
    if row_number in OTP:
        return "Approved OTP sandbox/test channel, synthetic person and consent"
    if row_number == 35:
        return "Explicit patient and visit write authorization; verify Save and Start Visit control"
    if row_number in PATIENT_WRITES:
        return "Explicit UAT write authorization, dedicated account, synthetic fixture and idempotency"
    if row_number in SAVE_VALIDATIONS:
        return "Authorized save attempt or isolated UI/API mock; verify no patient was created"
    if row_number == 32:
        return "Product decision: workbook expects Male/Female/Other dropdown; observed UAT form has Male/Female radios"
    if row_number == 38:
        return "Agreed device/network benchmark for workbook's <2-second threshold"
    return "Read-only UAT account and current form control confirmation"


def excel_safe(value: str) -> str:
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def next_action(row_number: int, progress: dict[str, str], duplicate: bool) -> str:
    if duplicate:
        return "Review exact duplicate with test owner; do not mark separately passed automatically"
    if row_number == 2:
        return "Keep passing read-only sidebar-to-Client-Registry regression"
    if row_number in {15, 23}:
        return "Extend beyond pre-submit scope before claiming full workbook case"
    if row_number == 20:
        return "Authorized live synthetic registration and UAT-assigned-ID reconciliation"
    if row_number == 3:
        return "Confirm National ID selector and implement read-only selection"
    if row_number in {4, 5}:
        return "Use approved synthetic/invalid ID fixture; verify validation before any external search"
    if row_number == 32:
        return "Resolve Male/Female/Other product mismatch; execute scoped radio check if authorized"
    if row_number == 38:
        return "Agree device/network benchmark, then implement form-load timing assertion"
    if row_number in OTP:
        return "Confirm OTP test channel and fixture before gated E2E implementation"
    if row_number in REGISTRY:
        return "Confirm synthetic HIE fixture and current UI before E2E implementation"
    if row_number in PATIENT_WRITES or row_number in SAVE_VALIDATIONS:
        return "Implement write-gated scenario only with approved synthetic data and cleanup/idempotency"
    if progress:
        return "Review scoped result before extending assertions"
    return "Confirm current UI; implement and execute read-only assertions"


def generate(workbook: Path, coverage_path: Path, progress_path: Path, output: Path) -> tuple[int, int]:
    sheet = load_workbook(workbook, read_only=True, data_only=True)[SHEET]
    rows = {n: next(sheet.iter_rows(min_row=n, max_row=n, values_only=True)) for n in range(2, 55)}
    coverage = read_csv_rows(coverage_path, "Source Key")
    progress = read_csv_rows(progress_path, "Source Key")
    records = []
    duplicates = 0

    for row_number in range(2, 55):
        source_key = PREFIX + str(row_number)
        values = rows[row_number]
        if not values[0] or not values[3] or source_key not in coverage:
            raise ValueError(f"Registration source row {row_number} is missing from workbook or inventory")
        canonical_row = row_number
        duplicate = row_number >= 39 and values[:9] == rows[row_number - 16][:9]
        if row_number >= 39 and not duplicate:
            raise ValueError(f"Expected repeat at {source_key} differs; review before regenerating")
        if duplicate:
            canonical_row -= 16
            duplicates += 1

        state = progress.get(source_key, {})
        if row_number == 20 and not state:
            state = {
                "Implementation State": "Partial implementation (write gated)",
                "Execution State": "Not executed on UAT",
                "Result": "Not run",
            }
        records.append({
            "Source Key": source_key,
            "Workbook TC ID": str(values[0]),
            "Canonical Source Key": PREFIX + str(canonical_row),
            "Exact Duplicate": "Yes" if duplicate else "No",
            "Workflow": workflow(canonical_row),
            "Scenario": str(values[3]).strip(),
            "Execution Safety": coverage[source_key]["Execution Safety"],
            "Implementation State": state.get("Implementation State", "Not implemented"),
            "Execution State": state.get("Execution State", "Not executed"),
            "Result": state.get("Result", "Not assessed"),
            "Required Decision or Data": dependency(canonical_row),
            "Next Action": next_action(canonical_row, state, duplicate),
        })

    if len(records) != 53 or duplicates != 16:
        raise ValueError("Expected 53 source rows with 16 exact repeated scenarios")
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", newline="", encoding="utf-8") as destination:
        writer = csv.DictWriter(destination, fieldnames=FIELDS, lineterminator="\n")
        writer.writeheader()
        for record in records:
            writer.writerow({field: excel_safe(record[field]) for field in FIELDS})
    return len(records), duplicates


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=PROJECT_ROOT / "docs/TaifaCare_Complete_Test_Cases_299_20251111.xlsx")
    parser.add_argument("--coverage", type=Path, default=PROJECT_ROOT / "docs/taifacare-test-case-coverage.csv")
    parser.add_argument("--progress", type=Path, default=PROJECT_ROOT / "docs/qa-automation-progress.csv")
    parser.add_argument("--output", type=Path, default=PROJECT_ROOT / "docs/registration-workflow-readiness.csv")
    args = parser.parse_args()
    count, duplicates = generate(args.workbook, args.coverage, args.progress, args.output)
    print(f"Wrote {count} source-keyed Registration rows; {duplicates} are exact repeats")


if __name__ == "__main__":
    main()
